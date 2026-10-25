#!/usr/bin/env python3
"""ייצוא מצגת HTML ל-PowerPoint ול-PDF.

PowerPoint: כל שלב בילדאפ הופך לשקף (צילום מלא של השקף). וידאו מקומי מוטמע מעליו
באותו מיקום ובאותו גודל, הערות הדובר נכנסות כהערות, ולינקים נשארים לחיצים.
PDF: הדפסה של המצגת דרך Chrome, שקף בכל עמוד, עם לינקים לחיצים.

שימוש:  python3 export.py <קובץ html> [--flat] [--only pptx|pdf]
        --flat  שקף אחד לכל שקף (בלי שלבי בילדאפ), נוח לשליחה ולא להצגה
        --only  רק פורמט אחד
מצב ההתקדמות נכתב ל-.work/export.json ליד המצגת.
"""
import json, sys, tempfile, time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import chromekit  # noqa: E402

EMU_PER_PX = 12192000 // chromekit.W


def px(v):
    return Emu(int(round(v * EMU_PER_PX)))


def status(deck, **kw):
    d = deck.parent / ".work"
    d.mkdir(exist_ok=True)
    (d / "export.json").write_text(json.dumps(dict(kw, time=time.strftime("%H:%M:%S")), ensure_ascii=False), "utf-8")


def add_link(slide, link, shot, out_png):
    """לינק לחיץ ב-PowerPoint: חותכים מצילום השקף בדיוק את אזור הכפתור, מניחים אותו
    במקום ושמים עליו את הלינק. נראה זהה, ו-PowerPoint למק מזהה עליו לחיצה
    (על צורה שקופה הוא מתעלם ומעביר שקף)."""
    x, y = int(link["x"]), int(link["y"])
    x2, y2 = int(round(link["x"] + link["w"])), int(round(link["y"] + link["h"]))
    Image.open(shot).crop((x, y, x2, y2)).save(out_png)
    pic = slide.shapes.add_picture(str(out_png), px(x), px(y), px(x2 - x), px(y2 - y))
    pic.click_action.hyperlink.address = link["href"]


def capture(chrome, tmp, base, si, k):
    """צילום של שקף אחד בשלב אחד, ומידע על הסרטונים והלינקים בו. כל צילום בפרופיל משלו כדי לרוץ במקביל."""
    url = f"{base}?export#{si + 1}.{k}"
    prof = tmp / f"prof_{si}_{k}"
    shot = tmp / f"s{si + 1}_{k}.png"
    chromekit.run(chrome, prof, f"--screenshot={shot}", url, budget=5000)
    info = chromekit.dom_json(chrome, prof, url, "export-info", budget=5000)["current"]
    if not shot.exists():
        raise RuntimeError(f"הצילום של שקף {si + 1} נכשל")
    return si, k, shot, info


def main():
    deck = Path(sys.argv[1]).resolve()
    flat = "--flat" in sys.argv
    only = sys.argv[sys.argv.index("--only") + 1] if "--only" in sys.argv else None
    want_pptx, want_pdf = only in (None, "pptx"), only in (None, "pdf")
    out_pptx = deck.with_suffix(".pptx")
    out_pdf = deck.with_suffix(".pdf")
    chrome = chromekit.find_chrome()
    if not chrome:
        status(deck, state="error", error="לא נמצא Chrome במחשב")
        sys.exit("לא נמצא דפדפן Chrome במחשב, ולכן אי אפשר לייצא.")
    base = deck.as_uri()
    t0 = time.time()
    status(deck, state="running")
    try:
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            slides = chromekit.dom_json(chrome, tmp / "p0", base + "?export#1", "export-info")["slides"] if want_pptx else []
            jobs = [(si, k) for si, sl in enumerate(slides) for k in ([sl["steps"]] if flat else range(sl["steps"] + 1))]
            with ThreadPoolExecutor(max_workers=6) as ex:
                shots = list(ex.map(lambda j: capture(chrome, tmp, base, *j), jobs))

            prs = Presentation()
            prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
            blank = prs.slide_layouts[6]
            for si, k, shot, info in shots:
                slide = prs.slides.add_slide(blank)
                jpg = shot.with_suffix(".jpg")  # JPEG במקום PNG: קובץ PowerPoint קטן פי כמה
                Image.open(shot).convert("RGB").save(jpg, quality=88)
                slide.shapes.add_picture(str(jpg), 0, 0, prs.slide_width, prs.slide_height)
                for i, v in enumerate(info["videos"]):
                    src = (deck.parent / (v["src"] or "")).resolve()
                    if not v["src"] or not src.exists():
                        continue
                    poster = tmp / f"p{si + 1}_{k}_{i}.png"
                    Image.open(shot).crop((int(v["x"]), int(v["y"]), int(v["x"] + v["w"]), int(v["y"] + v["h"]))).save(poster)
                    slide.shapes.add_movie(str(src), px(v["x"]), px(v["y"]), px(v["w"]), px(v["h"]),
                                           poster_frame_image=str(poster), mime_type="video/mp4")
                for j, link in enumerate(info["links"]):
                    add_link(slide, link, shot, tmp / f"l{si + 1}_{k}_{j}.png")
                notes = slides[si]["notes"].strip()
                if notes:
                    slide.notes_slide.notes_text_frame.text = notes
            if want_pptx:
                prs.save(out_pptx)
            if want_pdf:
                chromekit.run(chrome, tmp / "pdf", "--no-pdf-header-footer", f"--print-to-pdf={out_pdf}", base + "?print", budget=6000)
    except Exception as e:  # noqa: BLE001
        status(deck, state="error", error=str(e)[:500])
        raise
    res = {"state": "done", "pptx": str(out_pptx) if want_pptx and out_pptx.exists() else None,
           "pdf": str(out_pdf) if want_pdf and out_pdf.exists() else None,
           "seconds": round(time.time() - t0, 1)}
    status(deck, **res)
    print(json.dumps(res, ensure_ascii=False))


if __name__ == "__main__":
    main()
